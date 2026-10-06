# -*- coding: utf-8 -*-
"""Crop real screenshots and insert into personal report DOCX files."""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt
from PIL import Image

BASE = Path(__file__).resolve().parents[1]
SHOT_SRC = BASE / "docs" / "screenshots" / "android"
OUT_DIR = BASE / "docs" / "screenshots"
REPORT_DIR = BASE / "docs"
ROOT = BASE.parent

OUT_DIR.mkdir(parents=True, exist_ok=True)


def _is_brown(r: int, g: int, b: int) -> bool:
    return r > 70 and r < 170 and g > 35 and g < 120 and b < 100 and r >= g >= b


def crop_app(src: Path, dest: Path) -> Path:
    img = Image.open(src).convert("RGB")
    w, h = img.size
    px = img.load()

    # Splash / full-brown screens: keep as-is
    brown = 0
    sample = 0
    step_x = max(1, w // 30)
    step_y = max(1, h // 30)
    for y in range(0, h, step_y):
        for x in range(0, w, step_x):
            sample += 1
            if _is_brown(*px[x, y]):
                brown += 1
    if sample and brown / sample > 0.6:
        img.save(dest, "PNG")
        return dest

    # Remove solid brown right panel
    best = w
    for x in range(int(w * 0.35), w):
        bcount = 0
        s = 0
        for y in range(0, h, max(1, h // 40)):
            s += 1
            if _is_brown(*px[x, y]):
                bcount += 1
        if s and bcount / s > 0.75:
            best = x
            break
    if best < w * 0.95:
        img = img.crop((0, 0, best, h))
    img.save(dest, "PNG")
    return dest


FILES = {
    "splash": "splash.png",
    "home": "home.png",
    "favorites": "favorites.png",
    "reviews": "reviews.png",
    "search": "search.png",
    "settings": "settings.png",
    "detail": "detail.png",
    "form": "form.png",
}


def prepare_shots() -> dict[str, Path]:
    prepared = {}
    for key, name in FILES.items():
        src = SHOT_SRC / name
        if not src.exists():
            print("MISSING", src)
            continue
        dest = OUT_DIR / f"{key}.png"
        crop_app(src, dest)
        prepared[key] = dest
        print("prepared", key, dest, dest.stat().st_size)
    return prepared


def add_shot(doc: Document, title: str, path: Path | None, desc: str, widgets: str, libs: str, state: str, data: str):
    doc.add_heading(title, level=1)
    doc.add_paragraph(f"Miêu tả chức năng/giao diện: {desc}")
    p = doc.add_paragraph("Ảnh chức năng/giao diện:")
    if path and path.exists():
        doc.add_picture(str(path), width=Cm(9.5))
        last = doc.paragraphs[-1]
        last.alignment = WD_ALIGN_PARAGRAPH.CENTER
    else:
        doc.add_paragraph("(Chưa có ảnh)")
    doc.add_paragraph("Chi tiết cài đặt:")
    doc.add_paragraph(f"- Widget sử dụng: {widgets}")
    doc.add_paragraph(f"- Thư viện/plugin: {libs}")
    doc.add_paragraph(f"- Quản lý trạng thái / kiến trúc: {state}")
    doc.add_paragraph(f"- Đọc/lưu trữ dữ liệu: {data}")


def build_report(out_path: Path, mssv: str, name: str, parts: list[dict], shots: dict[str, Path]):
    d = Document()
    d.add_heading("BÁO CÁO DỰ ÁN CUỐI KỲ", level=0)
    d.add_paragraph("HỌC KỲ 1, NĂM HỌC 2026-2027")
    d.add_paragraph("CT484: PHÁT TRIỂN ỨNG DỤNG DI ĐỘNG")
    d.add_paragraph("Tên dự án/ứng dụng: CafeSpot – Ứng dụng review quán cà phê")
    d.add_paragraph(
        "Link GitHub mã nguồn: https://github.com/vuongdc24v7x414/CT484_CafeSpot_ThongVuong"
    )
    d.add_paragraph(f"MSSV: {mssv}")
    d.add_paragraph(f"Họ tên SV: {name}")
    d.add_paragraph("Lớp học phần: CT484")
    d.add_paragraph("")
    d.add_paragraph(
        "Miêu tả dự án/ứng dụng: CafeSpot là ứng dụng Flutter giúp người dùng "
        "khám phá, thêm/sửa/xóa quán cà phê, đánh giá và lưu yêu thích. Dữ liệu "
        "lưu cục bộ bằng SQLite (mobile), điều hướng bằng go_router, quản lý "
        "trạng thái bằng Provider, có thông báo cục bộ và giao diện responsive. "
        "Dự án thực hiện theo nhóm 2 thành viên."
    )
    d.add_paragraph(
        "(Sinh viên viết báo cáo cá nhân phần mình được phân công trong dự án)"
    )
    for i, part in enumerate(parts, 1):
        add_shot(
            d,
            f"Chức năng/giao diện {i}: {part['title']}",
            shots.get(part["shot"]),
            part["desc"],
            part["widgets"],
            part["libs"],
            part["state"],
            part["data"],
        )
    d.save(str(out_path))
    print("saved", out_path)


thong_parts = [
    dict(
        title="Splash Screen",
        shot="splash",
        desc="Màn hình khởi động hiển thị logo CafeSpot, tải dữ liệu rồi chuyển sang Home.",
        widgets="Scaffold, Column, Image.asset, CircularProgressIndicator, Text",
        libs="provider; go_router",
        state="Gọi load() của Provider trước khi điều hướng.",
        data="Kích hoạt đọc SharedPreferences + SQLite qua Provider.",
    ),
    dict(
        title="Trang chủ (Home)",
        shot="home",
        desc="Danh sách quán dạng Grid/List responsive, lọc danh mục, mở chi tiết hoặc thêm quán.",
        widgets="Scaffold, AppBar, GridView, FilterChip, FloatingActionButton, CafeCard, LayoutBuilder",
        libs="go_router, provider",
        state="CafeProvider (ChangeNotifier).",
        data="Đọc bảng cafes từ SQLite.",
    ),
    dict(
        title="Chi tiết quán",
        shot="detail",
        desc="Xem thông tin quán, yêu thích, sửa/xóa, gửi và xem đánh giá.",
        widgets="ListView, Image.network, RatingStars, TextField, AlertDialog, IconButton",
        libs="go_router, provider, intl, flutter_local_notifications",
        state="State cục bộ + CafeProvider CRUD.",
        data="Bảng cafes + reviews (SQLite).",
    ),
    dict(
        title="Form thêm/sửa quán",
        shot="form",
        desc="Form nhập tên, địa chỉ, mô tả, URL ảnh, danh mục; validate và lưu.",
        widgets="Form, TextFormField, DropdownButtonFormField, FilledButton",
        libs="provider, go_router, flutter_local_notifications",
        state="StatefulWidget controllers; CafeProvider.saveCafe.",
        data="INSERT/UPDATE bảng cafes.",
    ),
    dict(
        title="Điều hướng go_router + Shell",
        shot="home",
        desc="Route splash/home/favorites/reviews/settings/search/cafe/:id/cafe-form; fade-slide; NavigationBar/Rail.",
        widgets="StatefulShellRoute, NavigationBar, NavigationRail, CustomTransitionPage",
        libs="go_router",
        state="Shell giữ tab index; dữ liệu từ Provider.",
        data="Truyền pathParameters/extra khi điều hướng.",
    ),
    dict(
        title="Tầng dữ liệu SQLite",
        shot="home",
        desc="DatabaseHelper tạo schema, seed dữ liệu; CafeRepository CRUD và tính lại rating.",
        widgets="Không UI riêng",
        libs="sqflite, path",
        state="Repository được CafeProvider gọi; notifyListeners.",
        data="File cafe_spot.db; bảng cafes, reviews.",
    ),
]

vuong_parts = [
    dict(
        title="Yêu thích (Favorites)",
        shot="favorites",
        desc="Hiển thị quán yêu thích dạng GridView; bỏ yêu thích hoặc mở chi tiết.",
        widgets="GridView, CafeCard, AppBar, LayoutBuilder",
        libs="provider, go_router",
        state="CafeProvider.favorites",
        data="SELECT cafes WHERE is_favorite=1.",
    ),
    dict(
        title="Đánh giá của tôi",
        shot="reviews",
        desc="ListView toàn bộ review; mở quán hoặc xóa review.",
        widgets="ListView.separated, Card, ListTile, RatingStars",
        libs="provider, go_router, intl",
        state="CafeProvider.myReviews",
        data="Đọc/xóa bảng reviews.",
    ),
    dict(
        title="Tìm kiếm",
        shot="search",
        desc="Tìm theo tên/địa chỉ/mô tả; kết quả ListView.",
        widgets="TextField, ListView, ListTile, CircleAvatar, RatingStars",
        libs="go_router, provider",
        state="CafeProvider.setFilter(query:)",
        data="Truy vấn LIKE trên SQLite.",
    ),
    dict(
        title="Cài đặt",
        shot="settings",
        desc="Đổi tên hiển thị, dark mode, bật/tắt thông báo, gửi notification thử.",
        widgets="SwitchListTile, TextField, ListTile, OutlinedButton",
        libs="shared_preferences, flutter_local_notifications, provider",
        state="SettingsProvider + SharedPreferences.",
        data="Key: dark_mode, notifications_enabled, display_name.",
    ),
    dict(
        title="Local Notifications",
        shot="settings",
        desc="Khởi tạo plugin, kênh CafeSpot; thông báo khi thêm quán/đánh giá hoặc nút thử.",
        widgets="Gọi từ Settings/Detail/Form",
        libs="flutter_local_notifications",
        state="Phụ thuộc SettingsProvider.notificationsEnabled",
        data="Không lưu nội dung notification.",
    ),
    dict(
        title="Responsive layout & Shell",
        shot="home",
        desc="NavigationBar (mobile) hoặc NavigationRail (rộng); ResponsiveCenter.",
        widgets="NavigationBar, NavigationRail, LayoutBuilder",
        libs="go_router StatefulShellRoute",
        state="navigationShell.currentIndex",
        data="Không.",
    ),
]


if __name__ == "__main__":
    shots = prepare_shots()
    build_report(
        REPORT_DIR / "Baocao_BK24V7X703_LuuMinhThong.docx",
        "BK24V7X703",
        "Lưu Minh Thông",
        thong_parts,
        shots,
    )
    build_report(
        REPORT_DIR / "Baocao_BK24V7X414_DoChiVuong.docx",
        "BK24V7X414",
        "Đỗ Chí Vương",
        vuong_parts,
        shots,
    )
    # Copy to assignment root
    for name in [
        "Baocao_BK24V7X703_LuuMinhThong.docx",
        "Baocao_BK24V7X414_DoChiVuong.docx",
    ]:
        src = REPORT_DIR / name
        dst = ROOT / name
        dst.write_bytes(src.read_bytes())
        print("copied", dst)
